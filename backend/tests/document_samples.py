"""测试用的小文档:真能被各自的库读出来的最小文件。PDF 和 docx 手工拼(不为测试多装写库),pptx / xlsx 用读的那个库写。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path


def pdf_bytes(pages: list[str]) -> bytes:
    """几页、每页一行字的 PDF(Helvetica,只放 ASCII)。xref 偏移按字节算。"""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + index * 2} 0 R" for index in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + len(pages) * 2
    for index, text in enumerate(pages):
        content_id = 4 + index * 2
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def docx_bytes(paragraphs: list[tuple[str, str]]) -> bytes:
    """(样式, 文字) 一串段落的 docx。样式 `Heading1` 在 styles.xml 里登记成「heading 1」,mammoth 认成 h1。"""
    body = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'
        if style else f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'
        for style, text in paragraphs
    )
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/_rels/document.xml.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>',
        "word/styles.xml": f'<?xml version="1.0" encoding="UTF-8"?><w:styles {ns}>'
            '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>',
        "word/document.xml": f'<?xml version="1.0" encoding="UTF-8"?><w:document {ns}><w:body>{body}</w:body></w:document>',
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, text in files.items():
            archive.writestr(name, text)
    return out.getvalue()


def png_bytes(color: str = "red") -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (32, 24), color).save(out, format="PNG")
    return out.getvalue()


def pptx_bytes(slides: list[tuple[str, list[str], str]], *, picture_on: int | None = None) -> bytes:
    """(标题, 要点, 备注) 一串幻灯片。`picture_on`:第几张(1 起)上放一张小图。"""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for number, (title, bullets, notes) in enumerate(slides, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if picture_on == number:
            slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(6), Inches(5), Inches(1))
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def epub_bytes(chapters: list[tuple[str, str]]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                         '<rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
        items = "".join(f'<item id="c{i}" href="c{i}.xhtml"/>' for i in range(len(chapters)))
        spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
        archive.writestr("OEBPS/content.opf", f'<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf"><manifest>{items}</manifest><spine>{spine}</spine></package>')
        for i, (title, text) in enumerate(chapters):
            archive.writestr(f"OEBPS/c{i}.xhtml", f"<html><body><h1>{title}</h1><p>{text}</p></body></html>")
    return out.getvalue()


def write(tmp: Path, name: str, data: bytes | str) -> Path:
    path = tmp / name
    path.write_bytes(data.encode("utf-8") if isinstance(data, str) else data)
    return path
