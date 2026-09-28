"""本地解析(ADR 0031 §3):一份文档 → 按页 / 按幻灯片 / 按表 / 按章切好的 Markdown,外加插图和页面图。

**只用纯 Python 库**,不出本机、不花钱:PDF 用 pdfium 抽文字、渲每页图;docx 用 mammoth 转 HTML 再转 Markdown;
pptx、xlsx 用 python-pptx、openpyxl;csv / txt / md / html / epub 用标准库和 markdownify。

**本机装了 LibreOffice 就多做两件事**(见 office):Office 文档转成 PDF 渲页面图(智能体能看到 PPT 的版式),
老格式(doc / ppt / xls)先转成新格式再解析。没装就只有文字和插图,老格式说清楚要装它。

管得了原生数字文档,管不了扫描件和复杂版式:PDF 抽出来的字很少时,结果上带一句提醒(可以换 MinerU 这类插件解析)。
"""

from __future__ import annotations

import csv
import io
import logging
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree

from app.core.i18n import LocalizedError

logger = logging.getLogger(__name__)

#: 页面图的长边(像素)。给视觉模型看版式够了,一页几百 KB。
PAGE_IMAGE_LONG_EDGE = 1400
#: 最多渲多少页的页面图 —— 几百页的 PDF 全渲一遍要几分钟、几百 MB,文字照样全抽。
MAX_PAGE_IMAGES = 200
#: 表格最多列多少行、多少列(一张十万行的表整张塞进 Markdown 没人读得完,也塞不进上下文)。
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 40
#: 一页平均少于这么多字,多半是扫描件 / 图片版 PDF。
SCANNED_CHARS_PER_PAGE = 30

Progress = Callable[[float, str], None]


class DocumentParseError(LocalizedError, ValueError):
    """这份文档本地解析不了(文件坏了、加密、老格式没装 LibreOffice)。"""

    status = 422


@dataclass
class Section:
    """切出来的一段:PDF 的一页、PPT 的一张幻灯片、表格的一张表、docx / epub 的一章。"""

    index: int
    title: str
    markdown: str
    #: 这一段的页面图(相对解析目录),没有是 None。
    image: str | None = None


@dataclass
class Parsed:
    #: 这份文档按什么切:page / slide / sheet / section。
    unit: str
    sections: list[Section]
    #: 抽出来的插图(相对解析目录)。
    images: list[str] = field(default_factory=list)
    #: 给人看的提醒(i18n key):扫描件、表格截断、没装 LibreOffice 所以没有页面图……
    notes: list[str] = field(default_factory=list)
    #: 按页的页面图,和段落对不上时单独记(docx 按章切,页面图按页):按页取时用。对得上的挂在各段上。
    page_images: list[str] = field(default_factory=list)

    @property
    def markdown(self) -> str:
        """全文。每段前一个 `<!-- page: N -->` 标记 —— 读的人按页引用、按页取都靠它;渲染时不显示。"""
        return "\n\n".join(f"<!-- {self.unit}: {one.index} -->\n{one.markdown.strip()}" for one in self.sections).strip() + "\n"


def _noop(_fraction: float, _message: str) -> None:
    return None


def parse_local(source: Path, out_dir: Path, *, on_progress: Progress = _noop) -> Parsed:
    """解析一份文档,插图写进 `out_dir/images/`、页面图写进 `out_dir/pages/`。按扩展名分派。"""
    from app.domain.documents import office

    (out_dir / "images").mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    if suffix in office.LEGACY_UPGRADES:
        #: 老格式先转成新格式 —— 没有纯 Python 的可靠读法。
        upgraded = office.convert(source, office.LEGACY_UPGRADES[suffix], out_dir / "_converted")
        if upgraded is None:
            raise DocumentParseError("docErr_legacyNeedsLibreOffice", ext=suffix.lstrip("."))
        source, suffix = upgraded, upgraded.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise DocumentParseError("docErr_formatNotSupported", ext=suffix.lstrip("."))
    try:
        parsed = parser(source, out_dir, on_progress)
    except DocumentParseError:
        raise
    except Exception as exc:  # noqa: BLE001 —— 坏文件、加密文件:库各有各的异常,对用户都是「读不了这份」
        logger.warning("本地解析 %s 失败", source.name, exc_info=True)
        raise DocumentParseError("docErr_unreadable", detail=str(exc)[:200]) from exc
    #: Office 文档的页面图:本机有 LibreOffice 就转 PDF 渲;没有就说一声。
    if suffix in (".docx", ".pptx") and not any(one.image for one in parsed.sections):
        pdf = office.convert(source, "pdf", out_dir / "_converted")
        if pdf is None:
            parsed.notes.append("docNote_noPageImages")
        else:
            images = _render_pdf_pages(pdf, out_dir, on_progress)
            if parsed.unit == "slide":
                for section, image in zip(parsed.sections, images):
                    section.image = image
            else:
                #: docx 按章切,和页对不上:页面图单独记,按页取时用。
                parsed.page_images = images
    return parsed


# ── PDF ─────────────────────────────────────────────────────────────────────────


def _render_pdf_pages(pdf_path: Path, out_dir: Path, on_progress: Progress) -> list[str]:
    import pypdfium2 as pdfium

    pages_dir = out_dir / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    document = pdfium.PdfDocument(str(pdf_path))
    try:
        rendered: list[str] = []
        count = min(len(document), MAX_PAGE_IMAGES)
        for index in range(count):
            page = document[index]
            width, height = page.get_size()
            scale = PAGE_IMAGE_LONG_EDGE / max(width, height, 1)
            image = page.render(scale=scale).to_pil()
            name = f"pages/{index + 1:03d}.png"
            image.save(out_dir / name, optimize=True)
            rendered.append(name)
            on_progress(0.5 + 0.5 * (index + 1) / max(count, 1), "docProgress_renderPages")
        return rendered
    finally:
        document.close()


def _parse_pdf(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    import pypdfium2 as pdfium

    try:
        document = pdfium.PdfDocument(str(source))
    except pdfium.PdfiumError as exc:
        #: 加了打开密码的 PDF 也走到这里。
        raise DocumentParseError("docErr_unreadable", detail=str(exc)[:200]) from exc
    try:
        sections: list[Section] = []
        total = len(document)
        for index in range(total):
            page = document[index]
            text = page.get_textpage().get_text_range().replace("\r\n", "\n").replace("\r", "\n")
            text = re.sub(r"[ \t]+\n", "\n", text).strip()
            title = next((line.strip() for line in text.splitlines() if line.strip()), "")
            sections.append(Section(index=index + 1, title=title[:80], markdown=text))
            on_progress(0.5 * (index + 1) / max(total, 1), "docProgress_readPages")
    finally:
        document.close()
    images = _render_pdf_pages(source, out_dir, on_progress)
    for section, image in zip(sections, images):
        section.image = image
    notes = []
    chars = sum(len(one.markdown) for one in sections)
    if sections and chars / len(sections) < SCANNED_CHARS_PER_PAGE:
        notes.append("docNote_littleText")
    if total > MAX_PAGE_IMAGES:
        notes.append("docNote_pageImagesCapped")
    return Parsed(unit="page", sections=sections, notes=notes)


# ── Word ────────────────────────────────────────────────────────────────────────


def _html_to_markdown(html: str) -> str:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    #: 脚本、样式整个拿掉 —— markdownify 的 strip 只去标签、留下里面的代码。
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    text = markdownify(str(soup), heading_style="ATX", bullets="-")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _split_by_headings(markdown: str) -> list[Section]:
    """按一级、二级标题切成几章;没有标题就是一整段。"""
    parts = re.split(r"(?m)^(?=#{1,2} )", markdown)
    sections = []
    for part in (one.strip() for one in parts):
        if not part:
            continue
        first = part.splitlines()[0]
        title = first.lstrip("#").strip() if first.startswith("#") else ""
        sections.append(Section(index=len(sections) + 1, title=title[:80], markdown=part))
    return sections or [Section(index=1, title="", markdown="")]


def _parse_docx(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    import mammoth

    saved: list[str] = []

    def keep_image(image) -> dict[str, str]:
        ext = (image.content_type or "image/png").split("/")[-1].replace("jpeg", "jpg").split("+")[0]
        name = f"images/{len(saved) + 1:03d}.{ext}"
        with image.open() as stream:
            (out_dir / name).write_bytes(stream.read())
        saved.append(name)
        return {"src": name}

    with source.open("rb") as handle:
        result = mammoth.convert_to_html(handle, convert_image=mammoth.images.img_element(keep_image))
    on_progress(0.5, "docProgress_readPages")
    return Parsed(unit="section", sections=_split_by_headings(_html_to_markdown(result.value)), images=saved)


# ── PowerPoint ──────────────────────────────────────────────────────────────────


def _table_markdown(rows: list[list[str]]) -> str:
    """一张表 → Markdown 表格。第一行当表头;单元格里的换行、竖线转义掉。"""
    if not rows:
        return ""
    width = max(len(row) for row in rows)

    def cell(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", " ").strip()

    padded = [[cell(value) for value in row] + [""] * (width - len(row)) for row in rows]
    head, body = padded[0], padded[1:]
    lines = ["| " + " | ".join(head) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)


def _parse_pptx(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(source))
    sections: list[Section] = []
    saved: list[str] = []
    slides = list(presentation.slides)

    def walk(shapes) -> list:
        """组合形状展开;按从上到下、从左到右排 —— 阅读顺序,不是插入顺序。"""
        flat = []
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                flat.extend(walk(shape.shapes))
            else:
                flat.append(shape)
        return sorted(flat, key=lambda one: ((one.top or 0) // 20, one.left or 0))

    for number, slide in enumerate(slides, start=1):
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        blocks: list[str] = [f"# {title}"] if title else []
        for shape in walk(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                paragraphs = [(paragraph.level, "".join(run.text for run in paragraph.runs).strip())
                              for paragraph in shape.text_frame.paragraphs]
                paragraphs = [(level, text) for level, text in paragraphs if text]
                if len(paragraphs) == 1 and paragraphs[0][0] == 0:
                    blocks.append(paragraphs[0][1])
                elif paragraphs:
                    #: 几段的文本框就是一串要点,按缩进级别排。
                    blocks.append("\n".join(f"{'  ' * level}- {text}" for level, text in paragraphs))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                blocks.append(_table_markdown(rows))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                except (AttributeError, ValueError):
                    continue  # 链接出去的图片、EMF 之类读不出字节
                name = f"images/{number:03d}-{len(saved) + 1:03d}.{image.ext}"
                (out_dir / name).write_bytes(image.blob)
                saved.append(name)
                blocks.append(f"![]({name})")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append("> " + notes.replace("\n", "\n> "))
        sections.append(Section(index=number, title=title[:80], markdown="\n\n".join(blocks)))
        on_progress(0.5 * number / max(len(slides), 1), "docProgress_readPages")
    return Parsed(unit="slide", sections=sections, images=saved)


# ── 表格 ────────────────────────────────────────────────────────────────────────


def _cap_rows(rows: list[list[str]], notes: list[str]) -> list[list[str]]:
    if len(rows) > MAX_TABLE_ROWS + 1 or any(len(row) > MAX_TABLE_COLS for row in rows):
        if "docNote_tableTruncated" not in notes:
            notes.append("docNote_tableTruncated")
    return [row[:MAX_TABLE_COLS] for row in rows[: MAX_TABLE_ROWS + 1]]


def _parse_xlsx(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    from openpyxl import load_workbook

    workbook = load_workbook(str(source), read_only=True, data_only=True)
    notes: list[str] = []
    sections: list[Section] = []
    try:
        sheets = list(workbook.worksheets)
        for number, sheet in enumerate(sheets, start=1):
            rows: list[list[str]] = []
            for values in sheet.iter_rows(values_only=True):
                if len(rows) > MAX_TABLE_ROWS + 1:
                    break
                row = ["" if value is None else str(value) for value in values]
                if any(row):
                    rows.append(row)
            while rows and all(not row[-1] for row in rows if row):
                rows = [row[:-1] for row in rows]
                if not any(rows):
                    break
            body = _table_markdown(_cap_rows(rows, notes)) if rows else ""
            sections.append(Section(index=number, title=sheet.title, markdown=f"# {sheet.title}\n\n{body}".strip()))
            on_progress(number / max(len(sheets), 1), "docProgress_readPages")
    finally:
        workbook.close()
    return Parsed(unit="sheet", sections=sections, notes=notes)


def _read_text(source: Path) -> str:
    """文本文件的编码:UTF-8(带不带 BOM)优先,读不通退到 GB18030 —— 国内 Windows 上存的 txt / csv 多半是它。"""
    raw = source.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_csv(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    text = _read_text(source)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    notes: list[str] = []
    rows = [row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)]
    return Parsed(unit="sheet", sections=[Section(index=1, title=source.stem, markdown=_table_markdown(_cap_rows(rows, notes)))],
                  notes=notes)


# ── 文本、网页、电子书 ──────────────────────────────────────────────────────────


def _parse_text(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    text = _read_text(source).replace("\r\n", "\n")
    if source.suffix.lower() in (".md", ".markdown"):
        return Parsed(unit="section", sections=_split_by_headings(text.strip()))
    return Parsed(unit="section", sections=[Section(index=1, title="", markdown=text.strip())])


def _parse_html(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text(source), "html.parser")
    body = soup.body or soup
    return Parsed(unit="section", sections=_split_by_headings(_html_to_markdown(str(body))))


def _parse_epub(source: Path, out_dir: Path, on_progress: Progress) -> Parsed:
    """EPUB 是一个 zip:container.xml 指向 OPF,OPF 的 spine 给出章节顺序。一章一段。"""
    from bs4 import BeautifulSoup

    with zipfile.ZipFile(source) as archive:
        container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
        rootfile = next(el.attrib["full-path"] for el in container.iter() if el.tag.endswith("rootfile"))
        opf = ElementTree.fromstring(archive.read(rootfile))
        base = str(Path(rootfile).parent)
        manifest = {el.attrib["id"]: el.attrib["href"] for el in opf.iter() if el.tag.endswith("item") and "id" in el.attrib}
        spine = [el.attrib["idref"] for el in opf.iter() if el.tag.endswith("itemref")]
        sections: list[Section] = []
        for number, idref in enumerate(spine, start=1):
            href = manifest.get(idref)
            if not href:
                continue
            path = f"{base}/{href}" if base not in ("", ".") else href
            try:
                html = archive.read(path).decode("utf-8", errors="replace")
            except KeyError:
                continue
            soup = BeautifulSoup(html, "html.parser")
            markdown = _html_to_markdown(str(soup.body or soup))
            if not markdown:
                continue
            heading = soup.find(["h1", "h2", "h3"])
            sections.append(Section(index=len(sections) + 1, title=(heading.get_text(strip=True) if heading else "")[:80],
                                    markdown=markdown))
            on_progress(number / max(len(spine), 1), "docProgress_readPages")
    return Parsed(unit="section", sections=sections or [Section(index=1, title="", markdown="")])


_PARSERS: dict[str, Callable[[Path, Path, Progress], Parsed]] = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
    ".xlsx": _parse_xlsx,
    ".csv": _parse_csv,
    ".md": _parse_text,
    ".markdown": _parse_text,
    ".txt": _parse_text,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".epub": _parse_epub,
}
